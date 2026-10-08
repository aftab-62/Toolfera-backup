from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_BREAK
import sys
r=Path(sys.argv[1] if len(sys.argv)>1 else '/workspace/scratch/utilityhub-mobile-qa11');r.mkdir(exist_ok=True)
font='/opt/codex/runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/share/fonts/truetype/LiberationSans-Regular.ttf'
for number in [1,2]:
 im=Image.new('RGB',(1275,1650),'white');d=ImageDraw.Draw(im);d.text((110,120),f'UtilityHub OCR verification page {number}',font=ImageFont.truetype(font,42),fill='#17243b')
 lines=['This document is a scanned page with editable English text.','The browser should recognize each page in sequence.','Project reference: UH-2026. Invoice total: 125.50 USD.','Please review the output before sharing it.','No file is uploaded to an external conversion service.']
 for i,line in enumerate(lines):d.text((110,230+i*65),line,font=ImageFont.truetype(font,28),fill='black')
 im.save(r/f'scan-{number}.png')
c=Canvas(str(r/'english-scanned.pdf'),pagesize=(612,792))
for n in [1,2]:c.drawImage(str(r/f'scan-{n}.png'),0,0,612,792);c.showPage()
c.save()
# 7.2 million decoded pixels repeated across four pages: old code rejected it.
im=Image.new('RGB',(3000,2400),'#e1eaff');d=ImageDraw.Draw(im)
for x in range(0,3000,30):d.line((x,0,3000-x,2400),fill=(x%200,90,180),width=4)
im.save(r/'large-photo.jpg',quality=85)
c=Canvas(str(r/'memory-heavy.pdf'),pagesize=(612,792))
for n in range(1,5):c.setFont('Helvetica-Bold',18);c.drawString(54,730,f'Memory strategy verification page {n}');c.setFont('Helvetica',11);c.drawString(54,700,'This text and high resolution image should convert without excessive image memory.');c.drawImage(str(r/'large-photo.jpg'),54,480,180,144);c.showPage()
c.save()
doc=Document();doc.add_heading('UtilityHub Word conversion verification',0)
p=doc.add_paragraph('An editable paragraph with ');p.add_run('bold text').bold=True;p.add_run(', italic text').italic=True;p.add_run(' and Unicode café — 2026.').font.color.rgb=RGBColor.from_string('2F5BFF')
doc.add_heading('Practical test content',1)
for item in ['First ordered item','Second ordered item']:doc.add_paragraph(item,style='List Number')
for item in ['One bullet','Another bullet']:doc.add_paragraph(item,style='List Bullet')
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Tool','Status','Count']):c.text=v
for values in [('Word to PDF','Ready','1'),('PDF OCR','English','2')]:
 for c,v in zip(t.add_row().cells,values):c.text=v
doc.add_picture(str(r/'scan-1.png'),width=Inches(2))
doc.add_page_break();doc.add_heading('Second page verification',1)
for i in range(1,12):doc.add_paragraph(f'Paragraph {i}: The PDF should keep selectable text and sensible page flow. A complex Microsoft Word layout may differ in this browser conversion.')
doc.save(r/'word-basic.docx')
(r/'invalid.docx').write_text('This is not a DOCX.')
(r/'legacy.doc').write_text('Unsupported legacy format test.')
print(r)
